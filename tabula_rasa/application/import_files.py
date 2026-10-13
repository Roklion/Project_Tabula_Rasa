"""Bounded file preparation for supported imports, without ledger writes."""

import csv
from dataclasses import dataclass
import io
import zipfile

from tabula_rasa.errors import StorageValidationError
from tabula_rasa.quickbooks_import import REQUIRED_COLUMNS, general_ledger_rows


MAX_IMPORT_BYTES = 25 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
CSV_REQUIRED_COLUMNS = {"Date", "Type", "Amount", "Account", "Category"}


@dataclass(frozen=True)
class PreparedImport:
    rows: list[dict[str, str]]
    columns: list[str]
    conversion_errors: list[str]


def validate_import_bytes(content: bytes) -> None:
    if not content or len(content) > MAX_IMPORT_BYTES:
        raise StorageValidationError("Upload a non-empty import file no larger than 25 MB.")


def read_csv_upload(content: bytes, *, required_columns=()) -> list[dict[str, str]]:
    """Keep source cells as text, validate headers and bound each submitted batch."""
    validate_import_bytes(content)
    try:
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        columns = next(reader, [])
        if (not columns or len(columns) > 100 or len(set(columns)) != len(columns)
                or any(not column.strip() for column in columns)):
            raise StorageValidationError("Use a CSV with unique, non-empty column headers.")
        missing = set(required_columns) - set(columns)
        if missing:
            raise StorageValidationError("Missing columns: " + ", ".join(sorted(missing)))
        rows = []
        for cells in reader:
            if not cells:
                continue
            if len(cells) > len(columns):
                raise StorageValidationError("A CSV row has more values than the column headers.")
            if len(rows) == MAX_IMPORT_ROWS:
                raise StorageValidationError("Import at most 1,000 rows per batch; split larger sources before submitting.")
            rows.append(dict(zip(columns, [*cells, *([""] * (len(columns) - len(cells)))], strict=True)))
        return rows
    except (UnicodeError, csv.Error):
        raise StorageValidationError("Could not read this CSV file. Check its format and UTF-8 encoding.") from None


def prepare_csv_import(content: bytes) -> PreparedImport:
    rows = read_csv_upload(content, required_columns=CSV_REQUIRED_COLUMNS)
    return PreparedImport(rows, list(rows[0]) if rows else sorted(CSV_REQUIRED_COLUMNS), [])


def _bounded_workbook(content: bytes) -> bytes:
    # XLSX itself is a ZIP. Bound expansion before the worksheet reader opens it.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise StorageValidationError("The QuickBooks workbook expands beyond the supported file size.")
    return content


def prepare_quickbooks_import(filename: str, content: bytes) -> PreparedImport:
    """Use the existing General Ledger conversion and retain conversion diagnostics."""
    validate_import_bytes(content)
    extension = filename.rsplit(".", 1)[-1].casefold()
    if extension == "csv":
        rows = read_csv_upload(content, required_columns=REQUIRED_COLUMNS)
        return PreparedImport(rows, list(rows[0]) if rows else sorted(REQUIRED_COLUMNS), [])
    if extension not in {"xlsx", "zip"}:
        raise StorageValidationError("Upload a QuickBooks CSV, XLSX, or supported ZIP export.")
    try:
        workbook_bytes = content
        if extension == "zip":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                workbooks = [item for item in archive.infolist()
                             if item.filename.casefold().endswith(".xlsx")
                             and not any(part.casefold() == "__macosx" for part in item.filename.split("/"))
                             and not item.is_dir()]
                if len(workbooks) != 1:
                    raise StorageValidationError("The ZIP must contain exactly one QuickBooks .xlsx workbook.")
                if workbooks[0].file_size > MAX_IMPORT_BYTES:
                    raise StorageValidationError("The QuickBooks workbook is larger than 25 MB.")
                workbook_bytes = archive.read(workbooks[0])
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(_bounded_workbook(workbook_bytes)), read_only=True, data_only=True)
        try:
            sheet = workbook["General Ledger"]
            if sheet.max_column and sheet.max_column > 100:
                raise StorageValidationError("The QuickBooks workbook has too many columns.")
            matrix = []
            for row in sheet.iter_rows(values_only=True):
                if len(matrix) == 10000:
                    raise StorageValidationError("The QuickBooks worksheet exceeds the supported size of 10,000 rows; split larger sources before submitting.")
                matrix.append(row)
            rows, errors = general_ledger_rows(matrix)
            if len(rows) > MAX_IMPORT_ROWS:
                raise StorageValidationError("Import at most 1,000 rows per batch; split larger sources before submitting.")
            return PreparedImport(rows, sorted(REQUIRED_COLUMNS), errors)
        finally:
            workbook.close()
    except StorageValidationError:
        raise
    except Exception:
        # Workbook readers can raise several XML/ZIP implementation exceptions.
        # None of their provider text may escape file preparation.
        raise StorageValidationError("Could not read this QuickBooks export. Check its format and file integrity.") from None
